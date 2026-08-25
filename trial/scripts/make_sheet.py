#!/usr/bin/env python
"""Build the formatted daily-report workbook.

Produces the layout the Google Sheet should have — KPI block, channel /
product-type / category / top-10 tables, and the movers — with the purple
header styling of the original report. Google converts the .xlsx to a Sheet on
upload and keeps the formatting, tabs and number formats.

    python scripts/make_sheet.py --date 2026-08-24

The Apps Script builds the same layout in place each morning; this exists so
the format can be created in one step without running anything.
"""

import argparse
import csv
import datetime as dt
import io
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))

PURPLE = "D9C7EC"
PURPLE_DK = "B9A3D4"
GREY = "F7F7FA"
TOTAL_BG = "EFEAF7"
GREEN = "127A2E"
RED = "C02626"
AMBER_BG = "FFF4E5"

THIN = Side(style="thin", color="D8D8E0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SECTIONS = [
    ("Channel Summary", "Channel Summary"),
    ("Product Type Summary", "Product Type Summary"),
    ("Category Summary - Publication Books", "Category Summary — Publication Books"),
    ("Top 10 Publication Books", "Top 10 Publication Books"),
    ("Top 10 Stationery & Merchandise", "Top 10 Stationery & Merchandise"),
]


def num(v):
    if v in ("", None):
        return None
    try:
        return float(v)
    except ValueError:
        return v


def pct(v):
    """Sheets stores percentages as fractions."""
    return None if v in ("", None) else float(v) / 100.0


def pick(rows, section, scope=None):
    return [r for r in rows
            if r["section"] == section and (scope is None or r["scope"] == scope)]


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=PURPLE)
        cell.font = Font(bold=True, size=10)
        cell.border = Border(left=Side(style="thin", color=PURPLE_DK),
                             right=Side(style="thin", color=PURPLE_DK),
                             top=Side(style="thin", color=PURPLE_DK),
                             bottom=Side(style="thin", color=PURPLE_DK))
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_table(ws, r, title, data, with_delta):
    ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=12)
    r += 1

    head = (["Name", "Orders", "Orders Change", "Revenue (Lakhs)", "Revenue Change",
             "AOV", "Revenue Contribution %"] if with_delta else
            ["Name", "Orders", "Revenue (Lakhs)", "AOV", "Revenue Contribution %"])
    for i, h in enumerate(head, 1):
        ws.cell(row=r, column=i, value=h)
    style_header(ws, r, len(head))
    ws.row_dimensions[r].height = 30
    r += 1

    for i, d in enumerate(data):
        is_total = d["entity"] == "TOTAL"
        vals = ([d["entity"], num(d["orders"]), pct(d["orders_change_pct"]),
                 num(d["revenue_lakhs"]), pct(d["revenue_change_pct"]),
                 num(d["aov"]), pct(d["revenue_contribution_pct"])] if with_delta else
                [d["entity"], num(d["orders"]), num(d["revenue_lakhs"]),
                 num(d["aov"]), pct(d["revenue_contribution_pct"])])

        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BOX
            if is_total:
                cell.fill = PatternFill("solid", fgColor=TOTAL_BG)
                cell.font = Font(bold=True)
            elif i % 2:
                cell.fill = PatternFill("solid", fgColor=GREY)

        if with_delta:
            ws.cell(row=r, column=2).number_format = "#,##0"
            ws.cell(row=r, column=3).number_format = "0.00%"
            ws.cell(row=r, column=4).number_format = "#,##0.00"
            ws.cell(row=r, column=5).number_format = "0.00%"
            ws.cell(row=r, column=6).number_format = "#,##0"
            ws.cell(row=r, column=7).number_format = "0.00%"
            for col, raw in ((3, d["orders_change_pct"]), (5, d["revenue_change_pct"])):
                if raw not in ("", None):
                    colour = GREEN if float(raw) >= 0 else RED
                    cur = ws.cell(row=r, column=col)
                    cur.font = Font(color=colour, bold=is_total)
        else:
            ws.cell(row=r, column=2).number_format = "#,##0"
            ws.cell(row=r, column=3).number_format = "#,##0.00"
            ws.cell(row=r, column=4).number_format = "#,##0"
            ws.cell(row=r, column=5).number_format = "0.00%"
        r += 1
    return r


def write_movers(ws, r, title, data):
    ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=12)
    r += 1
    head = ["Product", "Yesterday Qty", "Day Before Qty", "Qty Diff"]
    for i, h in enumerate(head, 1):
        ws.cell(row=r, column=i, value=h)
    style_header(ws, r, len(head))
    r += 1

    for i, d in enumerate(data):
        vals = [d["entity"], num(d["qty"]), num(d["qty_prev"]), num(d["qty_diff"])]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BOX
            if i % 2:
                cell.fill = PatternFill("solid", fgColor=GREY)
            if c > 1:
                cell.number_format = "#,##0"
        diff = num(d["qty_diff"]) or 0
        ws.cell(row=r, column=4).font = Font(color=GREEN if diff >= 0 else RED)
        r += 1
    return r


def build(rows, report_date):
    wb = Workbook()

    # ── Daily Report ─────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Daily Report"
    r = 1
    ws.cell(row=r, column=1, value="PW Store — Daily Report").font = Font(bold=True, size=16)
    r += 1
    d = dt.date.fromisoformat(report_date)
    ws.cell(row=r, column=1,
            value=d.strftime("%A, %d %B %Y") + "   (data is for the previous day)"
            ).font = Font(color="666666")
    r += 2

    # Completeness warning, if the day has not landed.
    comp = pick(rows, "Data Completeness")
    if comp and comp[0]["revenue_contribution_pct"] not in ("", None):
        done = float(comp[0]["revenue_contribution_pct"])
        if done < 90:
            msg = ("PROVISIONAL — this day is about %.0f%% complete (%s orders against a "
                   "typical settled day of %s). Orders keep arriving for about a week because "
                   "marketplace channels report late, so treat every figure as a floor and be "
                   "careful with the day-on-day percentages."
                   % (done, "{:,}".format(int(float(comp[0]["orders"]))),
                      "{:,}".format(int(float(comp[0]["aov"])))))
            cell = ws.cell(row=r, column=1, value=msg)
            cell.fill = PatternFill("solid", fgColor=AMBER_BG)
            cell.font = Font(color="8A4B08", bold=True, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=7)
            ws.row_dimensions[r].height = 30
            r += 3

    # KPI block
    y = (pick(rows, "Headline", "Yesterday") or [{}])[0]
    m = (pick(rows, "Headline", "MTD") or [{}])[0]
    for i, h in enumerate(["Metric", "Value", "Change vs day before"], 1):
        ws.cell(row=r, column=i, value=h)
    style_header(ws, r, 3)
    r += 1
    for label, val, chg, fmt in (
        ("Yesterday's Orders", num(y.get("orders")), pct(y.get("orders_change_pct")), "#,##0"),
        ("Yesterday's Revenue (Lakhs)", num(y.get("revenue_lakhs")), pct(y.get("revenue_change_pct")), "#,##0.00"),
        ("MTD Orders", num(m.get("orders")), None, "#,##0"),
        ("MTD Revenue (Lakhs)", num(m.get("revenue_lakhs")), None, "#,##0.00"),
    ):
        ws.cell(row=r, column=1, value=label).border = BOX
        c2 = ws.cell(row=r, column=2, value=val)
        c2.number_format = fmt
        c2.border = BOX
        c2.font = Font(bold=True, size=12)
        c3 = ws.cell(row=r, column=3, value=chg)
        c3.number_format = "0.00%"
        c3.border = BOX
        if chg is not None:
            c3.font = Font(color=GREEN if chg >= 0 else RED)
        r += 1
    r += 1

    for key, title in SECTIONS:
        for scope in ("Yesterday", "MTD"):
            data = pick(rows, key, scope)
            if not data:
                continue
            r = write_table(ws, r, "%s — %s" % (title, scope), data, scope == "Yesterday")
            r += 1

    for key, title in (("Inclined Products", "Inclined Products (vs day before)"),
                       ("Declined Products", "Declined Products (vs day before)")):
        data = pick(rows, key, "Yesterday")
        if data:
            r = write_movers(ws, r, title, data)
            r += 1

    ws.column_dimensions["A"].width = 62
    for c in range(2, 8):
        ws.column_dimensions[get_column_letter(c)].width = 17
    ws.freeze_panes = "A2"

    # ── Raw Data ─────────────────────────────────────────────────────────────
    raw = wb.create_sheet("Raw Data")
    cols = list(rows[0].keys())
    raw.append(cols)
    style_header(raw, 1, len(cols))
    for rec in rows:
        raw.append([num(rec[c]) if c not in ("report_date", "section", "scope", "entity")
                    else rec[c] for c in cols])
    raw.freeze_panes = "A2"
    raw.column_dimensions["B"].width = 34
    raw.column_dimensions["E"].width = 60
    for c in ("A", "C", "D", "F", "G", "H", "I", "J", "K", "L", "M", "N"):
        raw.column_dimensions[c].width = 15

    # ── How it works ─────────────────────────────────────────────────────────
    info = wb.create_sheet("How it works")
    for i, line in enumerate([
        ["PW Store Daily Report"],
        [],
        ["Where the numbers come from"],
        ["Built from the PW Orders Intelligence Hub, the same backend the dashboard uses,"],
        ["so this sheet and the dashboard can never disagree for the same date range."],
        [],
        ["How it refreshes"],
        ["1. 08:15  a scheduled task rebuilds yesterday's figures and publishes them to"],
        ["          https://pw-orders-live.vercel.app/data/daily_report.csv"],
        ["2. 09:00  the Apps Script in this sheet reads that CSV, rewrites the tabs,"],
        ["          and emails the report."],
        [],
        ["Read this before trusting a day"],
        ["Orders keep arriving for about a week - marketplaces report late. Measured on"],
        ["25 Aug 2026: the previous day held only 34% of a settled day's orders, D-3 held"],
        ["69%, and a day did not reach ~98% until D-7. The Daily Report tab shows an amber"],
        ["banner with the completeness estimate. Day-on-day percentages on a provisional"],
        ["day compare two part-finished days and should not be read as trading movements."],
        [],
        ["Not included yet"],
        ["Daily Active Users - comes from app analytics, not the orders data."],
        ["MTD Target / Achieved% - a planning input; supply targets and the columns appear."],
    ]):
        info.append(line)
    info["A1"].font = Font(bold=True, size=14)
    for row_i in (3, 7, 13, 20):
        info.cell(row=row_i, column=1).font = Font(bold=True, size=11)
    info.column_dimensions["A"].width = 95
    return wb


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date", required=True)
    ap.add_argument("--csv", help="CSV to read (defaults to scripts/daily_report_<date>.csv)")
    ap.add_argument("--out")
    a = ap.parse_args()

    src = a.csv or os.path.join(HERE, "daily_report_%s.csv" % a.date)
    if not os.path.exists(src):
        print("no CSV at %s - run daily_report.py --csv first" % src)
        return 1

    rows = list(csv.DictReader(io.open(src, encoding="utf-8")))
    wb = build(rows, a.date)
    out = a.out or os.path.join(HERE, "PW_Store_Daily_Report.xlsx")
    wb.save(out)
    print("wrote %s (%.0f KB, %d tabs)"
          % (out, os.path.getsize(out) / 1024.0, len(wb.sheetnames)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
